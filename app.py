"""
TRA EFD Invoice Verification Portal
Run: python app.py
"""
from __future__ import annotations

import os
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import qrcode
from docx import Document
from docx.shared import Inches
from flask import Flask, abort, flash, redirect, render_template, request, send_file, url_for
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_FOLDER = BASE_DIR / "uploads"
OUTPUT_FOLDER = BASE_DIR / "outputs"
QRCODE_FOLDER = BASE_DIR / "qrcodes"
DATABASE = Path(os.environ.get("PORTAL_DATABASE", str(BASE_DIR / "portal.db")))
ALLOWED_EXTENSIONS = {".docx"}
MAX_CONTENT_LENGTH = 16 * 1024 * 1024  # 16 MB

for folder in (UPLOAD_FOLDER, OUTPUT_FOLDER, QRCODE_FOLDER):
    folder.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-before-production")
app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT_LENGTH


def db_connect():
    conn = sqlite3.connect(DATABASE, timeout=15)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db_connect() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                filename TEXT NOT NULL,
                verification_code TEXT NOT NULL,
                verification_link TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_logs_created_at ON logs(created_at)")


def clean_text(value: str) -> str:
    return (value or "").strip()


def valid_verification_time(value: str) -> str | None:
    """Accept HH:MM:SS or HHMMSS and return HHMMSS."""
    value = clean_text(value)
    if re.fullmatch(r"\d{2}:\d{2}:\d{2}", value):
        parts = value.split(":")
        hh, mm, ss = map(int, parts)
        if hh < 24 and mm < 60 and ss < 60:
            return f"{hh:02d}{mm:02d}{ss:02d}"
        return None
    if re.fullmatch(r"\d{6}", value):
        hh, mm, ss = int(value[:2]), int(value[2:4]), int(value[4:])
        if hh < 24 and mm < 60 and ss < 60:
            return value
    return None


def replace_placeholder_in_paragraph(paragraph, code: str, link: str, qr_path: Path) -> bool:
    """Replace <<EFD_SECTION>> while preserving a usable paragraph and add a QR image."""
    if "<<EFD_SECTION>>" not in paragraph.text:
        return False
    # Clear paragraph content, retaining paragraph formatting.
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run(f"Verification Code: {code}")
    paragraph.add_run().add_break()
    paragraph.add_run(f"Verification Link: {link}")
    paragraph.add_run().add_break()
    paragraph.add_run().add_picture(str(qr_path), width=Inches(0.8))
    return True


def process_docx(source: Path, destination: Path, code: str, link: str, qr_path: Path):
    doc = Document(str(source))
    inserted = False

    # Body paragraphs and table-cell paragraphs.
    for paragraph in doc.paragraphs:
        if replace_placeholder_in_paragraph(paragraph, code, link, qr_path):
            inserted = True

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_placeholder_in_paragraph(paragraph, code, link, qr_path):
                        inserted = True

    # Headers and footers can contain the placeholder too.
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            for paragraph in part.paragraphs:
                if replace_placeholder_in_paragraph(paragraph, code, link, qr_path):
                    inserted = True
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            if replace_placeholder_in_paragraph(paragraph, code, link, qr_path):
                                inserted = True

    # If the placeholder is absent, append a clearly labelled section instead.
    if not inserted:
        doc.add_paragraph(f"Verification Code: {code}")
        doc.add_paragraph(f"Verification Link: {link}")
        doc.add_picture(str(qr_path), width=Inches(0.8))

    doc.save(str(destination))


init_db()


@app.errorhandler(413)
def too_large(_error):
    return render_template("error.html", message="File is too large. Maximum upload size is 16 MB."), 413


@app.route("/", methods=["GET"])
def home():
    with db_connect() as conn:
        total = conn.execute("SELECT COUNT(*) FROM logs").fetchone()[0]
        # Use local application date, matching the timestamp stored by this app.
        today = datetime.now().strftime("%Y-%m-%d")
        today_count = conn.execute(
            "SELECT COUNT(*) FROM logs WHERE substr(created_at, 1, 10) = ?", (today,)
        ).fetchone()[0]
        recent = conn.execute("SELECT * FROM logs ORDER BY id DESC LIMIT 5").fetchall()
    return render_template("index.html", total=total, today=today_count, recent=recent)


@app.route("/process", methods=["POST"])
def process():
    uploaded = request.files.get("invoice")
    code = clean_text(request.form.get("verification_code", ""))
    raw_time = request.form.get("verification_time", "")
    clean_time = valid_verification_time(raw_time)

    if not uploaded or not uploaded.filename:
        flash("Please choose a Word invoice (.docx) to process.", "error")
        return redirect(url_for("home"))
    original_name = secure_filename(uploaded.filename)
    if not original_name or Path(original_name).suffix.lower() not in ALLOWED_EXTENSIONS:
        flash("Only Microsoft Word .docx files are supported. Legacy .doc files are not supported.", "error")
        return redirect(url_for("home"))
    if not code:
        flash("Enter the TRA verification code.", "error")
        return redirect(url_for("home"))
    if len(code) > 100 or any(ch.isspace() for ch in code):
        flash("The verification code must be 1–100 characters with no spaces.", "error")
        return redirect(url_for("home"))
    if not clean_time:
        flash("Enter a valid receipt time in HH:MM:SS or HHMMSS format.", "error")
        return redirect(url_for("home"))

    # Unique internal names prevent path traversal and avoid overwriting files.
    job_id = uuid4().hex
    source_path = UPLOAD_FOLDER / f"{job_id}_{original_name}"
    qr_path = QRCODE_FOLDER / f"{job_id}.png"
    output_name = f"PROCESSED_{Path(original_name).stem}_{job_id[:8]}.docx"
    output_path = OUTPUT_FOLDER / output_name
    link = f"https://verify.tra.go.tz/{code}_{clean_time}"

    try:
        uploaded.save(source_path)
        qrcode.make(link).save(qr_path)
        process_docx(source_path, output_path, code, link, qr_path)

        with db_connect() as conn:
            conn.execute(
                "INSERT INTO logs(filename, verification_code, verification_link, created_at) VALUES (?, ?, ?, ?)",
                (original_name, code, link, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            )

        return send_file(
            output_path,
            as_attachment=True,
            download_name=output_name,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except Exception:
        app.logger.exception("Failed to process invoice")
        flash("The invoice could not be processed. Confirm it is a valid, unencrypted .docx file and try again.", "error")
        return redirect(url_for("home"))
    finally:
        # Retain output and QR for audit/re-download; remove temporary source.
        try:
            source_path.unlink(missing_ok=True)
        except Exception:
            app.logger.warning("Could not remove temporary upload: %s", source_path)


@app.route("/history", methods=["GET"])
def history():
    with db_connect() as conn:
        records = conn.execute("SELECT * FROM logs ORDER BY id DESC LIMIT 500").fetchall()
    return render_template("history.html", records=records)


@app.route("/health", methods=["GET"])
def health():
    return {"status": "ok"}, 200


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    # Debug is deliberately disabled by default.
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=port, debug=False)
